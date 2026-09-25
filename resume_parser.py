from io import BytesIO

from docx import Document
from pypdf import PdfReader


MAX_FILE_SIZE = 5 * 1024 * 1024  # 5 MB


def extract_resume_text(uploaded_file):
    filename = uploaded_file.name.lower()
    file_bytes = uploaded_file.getvalue()

    if len(file_bytes) > MAX_FILE_SIZE:
        raise ValueError("File is too large. Please upload a file under 5 MB.")

    if filename.endswith(".pdf"):
        try:
            reader = PdfReader(BytesIO(file_bytes))
            text = "\n".join(page.extract_text() or "" for page in reader.pages)
        except Exception as exc:
            raise ValueError("Could not read this PDF. Try another PDF.") from exc

    elif filename.endswith(".docx"):
        try:
            document = Document(BytesIO(file_bytes))
            paragraphs = [p.text for p in document.paragraphs]

            for table in document.tables:
                for row in table.rows:
                    paragraphs.extend(cell.text for cell in row.cells)

            text = "\n".join(paragraphs)
        except Exception as exc:
            raise ValueError("Could not read this DOCX file.") from exc

    else:
        raise ValueError("Only PDF and DOCX files are accepted.")

    if not text.strip():
        raise ValueError(
            "No text was found. If this is a scanned PDF, use a text-based resume."
        )

    return text